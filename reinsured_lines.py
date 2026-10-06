from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


REQUIRED_WORKSHEET = "Reinsured Lines"
REQUIRED_COLUMNS = {
    "SubjectType",
    "Subject",
    "Domicile",
    "Assured",
    "Inception",
    "Expiry",
    "Attachment",
    "AttachmentIsFixed",
    "CoverType",
    "IsInitial",
    "Ccy",
    "Limit",
    "Excess",
    "AggregateLimit",
    "AggregateExcess",
    "Deductible",
    "Franchise",
    "CedablePercent",
    "Premium",
    "LinePer",
    "InterestPercent",
    "PolicyNo",
    "Broker",
    "BrokeragePercent",
    "Status",
    "Comments",
}

STRING_LIMITS = {
    "SubjectType": 12,
    "Subject": 150,
    "Domicile": 40,
    "Assured": 150,
    "CoverType": 12,
    "Ccy": 12,
    "Broker": 12,
    "Comments": 2000,
    "PolicyNo": 40,
}

NUMERIC_RULES: dict[str, tuple[bool, Decimal | None, Decimal | None, Decimal | None]] = {
    "Limit": (True, None, Decimal(0), None),
    "Excess": (True, Decimal(0), Decimal(0), None),
    "AggregateLimit": (False, None, Decimal(0), None),
    "AggregateExcess": (True, Decimal(0), Decimal(0), None),
    "Deductible": (True, Decimal(0), Decimal(0), None),
    "Franchise": (True, Decimal(0), Decimal(0), None),
    "CedablePercent": (True, Decimal(100), Decimal(0), Decimal(100)),
    "Premium": (True, Decimal(0), Decimal(0), None),
    "LinePer": (True, Decimal(100), Decimal(0), Decimal(100)),
    "InterestPercent": (True, Decimal(100), Decimal(0), Decimal(100)),
    "BrokeragePercent": (True, Decimal(0), Decimal(0), Decimal(100)),
}

DATE_FIELDS = ("Inception", "Expiry", "Attachment")
MONEY_FIELDS = {
    "Limit",
    "Excess",
    "AggregateLimit",
    "AggregateExcess",
    "Deductible",
    "Franchise",
    "Premium",
}


class WorkbookFormatError(ValueError):
    """Raised when an uploaded workbook does not match the ALPS import template."""


@dataclass
class ReinsuredWorkbook:
    reinsured: str
    as_at: Any
    rows: list[dict[str, Any]]
    worksheet: str


@dataclass
class RowValidation:
    excel_row: int
    values: dict[str, Any]
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    ignored: bool = False

    @property
    def can_import(self) -> bool:
        return not self.errors and not self.ignored


@dataclass
class ValidationSummary:
    rows: list[RowValidation]

    @property
    def error_count(self) -> int:
        return sum(bool(row.errors) for row in self.rows)

    @property
    def warning_count(self) -> int:
        return sum(bool(row.warnings) for row in self.rows)

    @property
    def ignored_count(self) -> int:
        return sum(row.ignored for row in self.rows)

    @property
    def import_count(self) -> int:
        return sum(row.can_import for row in self.rows)

    @property
    def valid(self) -> bool:
        return bool(self.rows) and self.error_count == 0


def validate_metadata(reinsured: str, as_at: Any) -> tuple[list[str], date | None]:
    errors: list[str] = []
    normalized_as_at: date | None = None

    if not reinsured.strip() or reinsured.strip().startswith("<"):
        errors.append("Enter the Reinsured name in the workbook.")
    elif len(reinsured.strip()) > 40:
        errors.append("Reinsured name exceeds 40 characters.")

    try:
        normalized_as_at = _parse_date(as_at)
    except ValueError:
        errors.append("AsAt must be a valid date.")
    if normalized_as_at is None:
        errors.append("Enter the AsAt date in the workbook.")
    return errors, normalized_as_at


def _named_range_cell(workbook: Any, name: str) -> Any:
    defined_name = workbook.defined_names.get(name)
    if defined_name is None:
        raise WorkbookFormatError(f"The workbook is missing the '{name}' named range.")

    try:
        sheet_name, cell_range = next(iter(defined_name.destinations))
    except (StopIteration, TypeError, ValueError) as exc:
        raise WorkbookFormatError(
            f"The '{name}' named range does not point to a worksheet cell."
        ) from exc

    if sheet_name not in workbook.sheetnames:
        raise WorkbookFormatError(
            f"The '{name}' named range refers to a missing worksheet."
        )

    min_col, min_row, _, _ = range_boundaries(cell_range)
    return workbook[sheet_name].cell(row=min_row + 1, column=min_col).value


def read_generic_workbook(file_bytes: bytes) -> ReinsuredWorkbook:
    """Read a macro-enabled or standard ALPS Generic Reinsured Lines workbook."""
    try:
        workbook = load_workbook(
            BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise WorkbookFormatError(
            "The file could not be opened as an Excel .xlsx or .xlsm workbook."
        ) from exc

    try:
        if REQUIRED_WORKSHEET not in workbook.sheetnames:
            raise WorkbookFormatError(
                f"The workbook must contain a '{REQUIRED_WORKSHEET}' worksheet."
            )

        reinsured = _named_range_cell(workbook, "Reinsured")
        as_at = _named_range_cell(workbook, "AsAtDate")
        defined_name = workbook.defined_names.get("ReinsuredLines")
        if defined_name is None:
            raise WorkbookFormatError(
                "The workbook is missing the 'ReinsuredLines' named range."
            )
        try:
            sheet_name, cell_range = next(iter(defined_name.destinations))
        except (StopIteration, TypeError, ValueError) as exc:
            raise WorkbookFormatError(
                "The 'ReinsuredLines' named range does not point to a worksheet table."
            ) from exc
        if sheet_name != REQUIRED_WORKSHEET:
            raise WorkbookFormatError(
                f"The 'ReinsuredLines' named range must refer to '{REQUIRED_WORKSHEET}'."
            )

        min_col, header_row, max_col, range_end_row = range_boundaries(cell_range)
        worksheet = workbook[sheet_name]
        headers = [
            str(worksheet.cell(row=header_row, column=column).value or "").strip()
            for column in range(min_col, max_col + 1)
        ]
        missing_columns = sorted(REQUIRED_COLUMNS.difference(headers))
        if missing_columns:
            raise WorkbookFormatError(
                "The ReinsuredLines table is missing required columns: "
                + ", ".join(missing_columns)
            )

        column_count = len(headers)
        rows: list[dict[str, Any]] = []
        first_data_row = header_row + 1
        for excel_row, cells in enumerate(
            worksheet.iter_rows(
                min_row=first_data_row,
                max_row=max(range_end_row, first_data_row),
                min_col=min_col,
                max_col=min_col + column_count - 1,
                values_only=True,
            ),
            start=first_data_row,
        ):
            row = dict(zip(headers, cells))
            if any(value is not None and str(value).strip() for value in row.values()):
                rows.append({"_excel_row": excel_row, **row})

        return ReinsuredWorkbook(
            reinsured=str(reinsured).strip() if reinsured is not None else "",
            as_at=as_at,
            rows=rows,
            worksheet=sheet_name,
        )
    finally:
        workbook.close()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _parse_date(value: Any) -> date | None:
    if _is_blank(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        from openpyxl.utils.datetime import from_excel

        converted = from_excel(value)
        return converted.date() if isinstance(converted, datetime) else converted
    if isinstance(value, str):
        candidate = value.strip()
        try:
            return datetime.fromisoformat(candidate).date()
        except ValueError:
            for date_format in ("%d/%m/%Y", "%d-%b-%Y", "%m/%d/%Y"):
                try:
                    return datetime.strptime(candidate, date_format).date()
                except ValueError:
                    pass
    raise ValueError("must be a valid date")


def _parse_decimal(value: Any) -> Decimal | None:
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        raise ValueError("must be numeric")
    if isinstance(value, Decimal):
        if not value.is_finite():
            raise ValueError("must be a finite number")
        return value
    try:
        number = Decimal(str(value).strip().replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError("must be numeric") from exc
    if not number.is_finite():
        raise ValueError("must be a finite number")
    return number


def _parse_boolean(value: Any) -> bool:
    if _is_blank(value):
        return False
    if isinstance(value, bool):
        return value
    text = str(value).strip().casefold()
    if text.startswith(("t", "y", "1")):
        return True
    if text.startswith(("f", "n", "0")):
        return False
    raise ValueError("must start with T, F, Y, N, 1, or 0")


def _validate_row(
    source: dict[str, Any],
    line_of_business: str,
) -> RowValidation:
    values = {key: value for key, value in source.items() if key != "_excel_row"}
    result = RowValidation(
        excel_row=int(source.get("_excel_row", 0)),
        values=values,
    )

    for field_name in DATE_FIELDS:
        value = values.get(field_name)
        if _is_blank(value) and field_name in ("Inception", "Expiry"):
            result.errors.append(f"{field_name} is required.")
            continue
        try:
            parsed = _parse_date(value)
            if parsed is not None:
                values[field_name] = parsed
        except ValueError as exc:
            result.errors.append(f"{field_name} {exc}.")

    inception = values.get("Inception")
    expiry = values.get("Expiry")
    if isinstance(inception, date) and isinstance(expiry, date):
        if expiry < inception:
            result.errors.append("Expiry must be later than Inception.")
        elif expiry == inception and not bool(values.get("IsInitial")):
            result.warnings.append("Inception and Expiry are the same date.")
    if _is_blank(values.get("Attachment")) and isinstance(inception, date):
        values["Attachment"] = inception

    for field_name in ("AttachmentIsFixed", "IsInitial"):
        try:
            values[field_name] = _parse_boolean(values.get(field_name))
        except ValueError as exc:
            result.errors.append(f"{field_name} {exc}.")

    parsed_numbers: dict[str, Decimal | None] = {}
    for field_name in NUMERIC_RULES:
        try:
            parsed_numbers[field_name] = _parse_decimal(values.get(field_name))
        except ValueError as exc:
            result.errors.append(f"{field_name} {exc}.")
            parsed_numbers[field_name] = None
        else:
            if parsed_numbers[field_name] is not None and field_name in MONEY_FIELDS:
                parsed_numbers[field_name] = parsed_numbers[field_name].quantize(
                    Decimal("0.01")
                )

    for field_name in ("LinePer", "InterestPercent"):
        if parsed_numbers[field_name] is None:
            result.warnings.append(
                f"{field_name} is blank; this row will not be imported."
            )
            result.ignored = True

    line_percent = parsed_numbers["LinePer"]
    interest_percent = parsed_numbers["InterestPercent"]
    limit = parsed_numbers["Limit"]
    if (
        not result.ignored
        and line_percent is not None
        and interest_percent is not None
        and line_percent * interest_percent == 0
        and (limit is None or limit == 0)
    ):
        result.ignored = True
        result.warnings.append(
            "LinePer, InterestPercent, and Limit indicate a zero-value line; "
            "the legacy importer skips it."
        )

    if result.ignored:
        return result

    for field_name, max_length in STRING_LIMITS.items():
        value = values.get(field_name)
        if not _is_blank(value):
            text = str(value).strip()
            values[field_name] = text
            if len(text) > max_length:
                result.errors.append(
                    f"{field_name} exceeds {max_length} characters."
                )

    for field_name in ("Subject", "CoverType", "Ccy", "Status"):
        if _is_blank(values.get(field_name)):
            result.errors.append(f"{field_name} is required.")

    status = values.get("Status")
    if not _is_blank(status):
        status_code = str(status).strip().casefold()
        if status_code in ("s", "signed"):
            values["Status"] = "Signed"
        elif status_code in ("e", "estimated"):
            values["Status"] = "Estimated"
        else:
            result.errors.append("Status must be S, Signed, E, or Estimated.")

    for field_name, (mandatory, default, minimum, maximum) in NUMERIC_RULES.items():
        parsed = parsed_numbers[field_name]
        if parsed is None:
            if default is not None:
                values[field_name] = default
            elif mandatory:
                result.errors.append(f"{field_name} is required.")
            continue

        if minimum is not None:
            if field_name in ("Limit", "AggregateLimit") and parsed <= minimum:
                result.errors.append(f"{field_name} must be greater than 0.")
            elif parsed < minimum:
                result.errors.append(f"{field_name} must be at least {minimum}.")
        if maximum is not None and parsed > maximum:
            result.errors.append(f"{field_name} must be at most {maximum}.")
        values[field_name] = parsed

    cover_type = values.get("CoverType")
    if not _is_blank(cover_type):
        cover_types = [part.strip() for part in str(cover_type).split(",")]
        if line_of_business in ("Casualty", "Credit and Political Risks"):
            counts = Counter(code.casefold() for code in cover_types)
            duplicates = sorted(code for code, count in counts.items() if count > 1)
            if duplicates:
                result.errors.append(
                    "CoverType contains duplicate codes: " + ", ".join(duplicates) + "."
                )
            if any(not code for code in cover_types):
                result.errors.append("CoverType contains an empty code.")
        elif len(cover_types) > 1:
            result.errors.append(
                "CoverType must contain one code for this Line of Business."
            )

    return result


def validate_rows(
    rows: list[dict[str, Any]],
    line_of_business: str = "Generic / Other",
) -> ValidationSummary:
    """Run workbook-local checks; ALPS database lookups are intentionally omitted."""
    return ValidationSummary(
        rows=[
            _validate_row(row, line_of_business)
            for row in rows
        ]
    )
