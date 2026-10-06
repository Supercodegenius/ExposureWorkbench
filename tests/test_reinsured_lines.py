import unittest
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from reinsured_lines import (
    REQUIRED_COLUMNS,
    WorkbookFormatError,
    read_generic_workbook,
    validate_metadata,
    validate_rows,
)


def valid_row(**overrides):
    values = {
        "SubjectType": "Property",
        "Subject": "Example Subject",
        "Domicile": "UK",
        "Assured": "Example Assured",
        "Inception": date(2025, 1, 1),
        "Expiry": date(2026, 1, 1),
        "Attachment": None,
        "AttachmentIsFixed": "Y",
        "CoverType": "PROP",
        "IsInitial": "N",
        "Ccy": "GBP",
        "Limit": 100000,
        "Excess": None,
        "AggregateLimit": None,
        "AggregateExcess": None,
        "Deductible": None,
        "Franchise": None,
        "CedablePercent": None,
        "Premium": None,
        "LinePer": 100,
        "InterestPercent": 100,
        "PolicyNo": None,
        "Broker": None,
        "BrokeragePercent": None,
        "Status": "S",
        "Comments": None,
    }
    values.update(overrides)
    return values


def workbook_bytes():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Reinsured Lines"
    worksheet["A2"] = "Example Reinsured"
    worksheet["A4"] = date(2025, 1, 1)

    headers = sorted(REQUIRED_COLUMNS)
    for column, name in enumerate(headers, start=1):
        worksheet.cell(row=12, column=column, value=name)
    for column, name in enumerate(headers, start=1):
        worksheet.cell(
            row=13,
            column=column,
            value=valid_row().get(name),
        )
    workbook.defined_names.add(
        DefinedName("Reinsured", attr_text="'Reinsured Lines'!$A$1:$A$2")
    )
    workbook.defined_names.add(
        DefinedName("AsAtDate", attr_text="'Reinsured Lines'!$A$3:$A$4")
    )
    last_column = worksheet.cell(row=12, column=len(headers)).column_letter
    workbook.defined_names.add(
        DefinedName(
            "ReinsuredLines",
            attr_text=f"'Reinsured Lines'!$A$12:${last_column}$13",
        )
    )
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class ReinsuredLinesTests(unittest.TestCase):
    def test_reads_template_named_ranges_and_line_data(self):
        workbook = read_generic_workbook(workbook_bytes())

        self.assertEqual(workbook.reinsured, "Example Reinsured")
        self.assertEqual(workbook.as_at.date(), date(2025, 1, 1))
        self.assertEqual(len(workbook.rows), 1)
        self.assertEqual(workbook.rows[0]["_excel_row"], 13)

    def test_rejects_workbook_without_alps_named_ranges(self):
        with self.assertRaises(WorkbookFormatError):
            read_generic_workbook(b"not an Excel workbook")

    def test_validates_and_normalizes_generic_row(self):
        summary = validate_rows([valid_row()])
        row = summary.rows[0]

        self.assertTrue(summary.valid)
        self.assertTrue(row.can_import)
        self.assertEqual(row.values["Attachment"], date(2025, 1, 1))
        self.assertEqual(row.values["Status"], "Signed")
        self.assertEqual(str(row.values["Excess"]), "0")

    def test_rejects_reversed_dates_and_invalid_percentages(self):
        summary = validate_rows(
            [
                valid_row(
                    Inception=date(2026, 1, 1),
                    Expiry=date(2025, 1, 1),
                    LinePer=101,
                )
            ]
        )

        self.assertFalse(summary.valid)
        self.assertIn("Expiry must be later than Inception.", summary.rows[0].errors)
        self.assertIn("LinePer must be at most 100.", summary.rows[0].errors)

    def test_blank_line_percentage_is_skipped_with_warning(self):
        summary = validate_rows([valid_row(LinePer=None, Subject=None)])

        self.assertEqual(summary.error_count, 0)
        self.assertEqual(summary.ignored_count, 1)
        self.assertFalse(summary.rows[0].can_import)
        self.assertTrue(any("LinePer is blank" in warning for warning in summary.rows[0].warnings))

    def test_casualty_rejects_duplicate_cover_codes_without_case_sensitivity(self):
        summary = validate_rows(
            [valid_row(CoverType="PROP, prop")],
            line_of_business="Casualty",
        )

        self.assertTrue(any("duplicate codes" in error for error in summary.rows[0].errors))

    def test_generic_rejects_multiple_cover_codes(self):
        summary = validate_rows([valid_row(CoverType="PROP, FIRE")])

        self.assertTrue(
            any("one code" in error for error in summary.rows[0].errors)
        )

    def test_metadata_requires_reinsured_and_as_at(self):
        errors, normalized_date = validate_metadata("<Insert Reinsured Name>", None)

        self.assertIsNone(normalized_date)
        self.assertEqual(len(errors), 2)


if __name__ == "__main__":
    unittest.main()
